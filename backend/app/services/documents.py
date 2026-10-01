"""PDF/DOCX/text ingestion with optional Tesseract OCR fallback."""
from __future__ import annotations

import io
import re
import shutil
from dataclasses import dataclass
from pathlib import Path


@dataclass
class ExtractedDocument:
    text: str
    pages: list[dict[str, str | int]]
    warnings: list[str]
    extraction_method: str


def page_text_ranges(pages: list[dict[str, str | int]]) -> list[dict[str, int]]:
    """Build compact page-to-text offsets without duplicating extracted text in storage."""
    ranges: list[dict[str, int]] = []
    offset = 0
    for page in pages:
        text = str(page.get("text", ""))
        if not text:
            continue
        if ranges:
            offset += 2  # _extract_pdf joins nonempty page text with two newlines.
        start = offset
        offset += len(text)
        ranges.append({"page": int(page.get("page", len(ranges) + 1)), "start": start, "end": offset})
    return ranges


def extract_document(filename: str, content: bytes, ocr_language: str = "eng") -> ExtractedDocument:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(content, ocr_language)
    if suffix == ".docx":
        return _extract_docx(content)
    if suffix == ".txt":
        text = _clean(content.decode("utf-8", errors="replace"))
        return ExtractedDocument(text, [{"page": 1, "text": text}], [], "plain_text")
    if suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff"}:
        return _extract_image(content, ocr_language)
    raise ValueError("Unsupported file type. Upload PDF, DOCX, TXT, PNG, JPG, or TIFF.")


def _extract_pdf(content: bytes, ocr_language: str) -> ExtractedDocument:
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:
        raise RuntimeError("PDF extraction requires PyMuPDF. Install backend/requirements.txt.") from exc
    try:
        doc = fitz.open(stream=content, filetype="pdf")
    except Exception as exc:
        raise ValueError("The PDF is corrupt, encrypted, or unreadable.") from exc
    if doc.is_encrypted:
        raise ValueError("Password-protected PDFs are not supported. Upload an unlocked copy.")
    pages: list[dict[str, str | int]] = []
    warnings: list[str] = []
    did_ocr = False
    for index, page in enumerate(doc, start=1):
        text = page.get_text("text") or ""
        if len(re.sub(r"\s", "", text)) < 30:
            ocr_text = _ocr_page(page, ocr_language)
            if ocr_text.strip():
                text = ocr_text
                did_ocr = True
            elif not text.strip():
                warnings.append(f"Page {index} had no extractable text; OCR may be unavailable or unsuccessful.")
        pages.append({"page": index, "text": _clean(text)})
    doc.close()
    full = "\n\n".join(str(p["text"]) for p in pages if p["text"])
    if not full.strip():
        raise ValueError("No text could be extracted. Check that the PDF is readable and OCR is installed for scanned files.")
    return ExtractedDocument(full, pages, warnings, "pymupdf+tesseract" if did_ocr else "pymupdf")


def _ocr_page(page, language: str) -> str:
    if not shutil.which("tesseract"):
        return ""
    try:
        import pytesseract
        pix = page.get_pixmap(matrix=__import__("fitz").Matrix(2, 2), alpha=False)
        from PIL import Image
        return pytesseract.image_to_string(Image.open(io.BytesIO(pix.tobytes("png"))), lang=language)
    except Exception:
        return ""


def _extract_image(content: bytes, language: str) -> ExtractedDocument:
    if not shutil.which("tesseract"):
        raise RuntimeError("Image OCR requires the Tesseract executable and pytesseract. Install Tesseract and see README troubleshooting.")
    try:
        import pytesseract
        from PIL import Image
        image = Image.open(io.BytesIO(content))
        text = _clean(pytesseract.image_to_string(image, lang=language))
    except Exception as exc:
        raise ValueError("Image OCR failed. Check that the image is readable and Tesseract language data is installed.") from exc
    if not text:
        raise ValueError("OCR did not find readable text in the image.")
    return ExtractedDocument(text, [{"page": 1, "text": text}], [], "tesseract_image")


def _extract_docx(content: bytes) -> ExtractedDocument:
    try:
        from docx import Document
        doc = Document(io.BytesIO(content))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = _clean("\n".join(parts))
    except Exception as exc:
        raise ValueError("The DOCX file is corrupt or unsupported.") from exc
    if not text:
        raise ValueError("No text was found in the DOCX file.")
    return ExtractedDocument(text, [{"page": 1, "text": text}], [], "python_docx")


def _clean(text: str) -> str:
    text = text.replace("\x00", " ").replace("\r", "\n")
    text = re.sub(r"[\t\xa0 ]+", " ", text)
    text = re.sub(r"\n[ ]+", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
